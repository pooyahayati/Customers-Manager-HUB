import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO
from typing import cast

from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from customers_manager_hub.knowledge_storage import PDF_MEDIA_TYPE, XLSX_MEDIA_TYPE

MAX_PDF_PAGES = 500
MAX_EXTRACTED_CHARACTERS = 2_000_000
MAX_XLSX_ZIP_ENTRIES = 10_000
MAX_XLSX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_SOURCE_CHUNKS = 512
CHUNK_MAX_CHARACTERS = 1_200
CHUNK_OVERLAP_CHARACTERS = 150
_XLSX_UNIT_TARGET_CHARACTERS = 4_800
_WHITESPACE_RUN = re.compile(r"[ \t\f\v]+")
_EXCESS_NEWLINES = re.compile(r"\n{3,}")


class KnowledgeParseError(ValueError):
    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass(frozen=True, slots=True)
class ExtractionUnit:
    text: str
    provenance: dict[str, object]


@dataclass(frozen=True, slots=True)
class ChunkDraft:
    content: str
    provenance: dict[str, object]


def normalize_extracted_text(value: str) -> str:
    normalized = value.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    normalized = "\n".join(
        _WHITESPACE_RUN.sub(" ", line).strip() for line in normalized.split("\n")
    )
    return _EXCESS_NEWLINES.sub("\n\n", normalized).strip()


def _validate_total_characters(units: list[ExtractionUnit]) -> None:
    total = sum(len(unit.text) for unit in units)
    if total == 0:
        raise KnowledgeParseError("knowledge_no_extractable_text")
    if total > MAX_EXTRACTED_CHARACTERS:
        raise KnowledgeParseError("knowledge_extracted_text_too_large")


def parse_pdf(data: bytes) -> list[ExtractionUnit]:
    try:
        reader = PdfReader(BytesIO(data), strict=False)
    except (PdfReadError, ValueError, OSError) as exc:
        raise KnowledgeParseError("knowledge_pdf_invalid") from exc
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception as exc:  # pypdf raises multiple crypto/parser exception types
            raise KnowledgeParseError("knowledge_pdf_encrypted") from exc
        if unlocked == 0:
            raise KnowledgeParseError("knowledge_pdf_encrypted")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise KnowledgeParseError("knowledge_pdf_too_many_pages")

    units: list[ExtractionUnit] = []
    extracted_total = 0
    for page_index, page in enumerate(reader.pages, start=1):
        try:
            raw_text = page.extract_text() or ""
        except (PdfReadError, ValueError, OSError) as exc:
            raise KnowledgeParseError("knowledge_pdf_extract_failed") from exc
        text = normalize_extracted_text(raw_text)
        if not text:
            continue
        extracted_total += len(text)
        if extracted_total > MAX_EXTRACTED_CHARACTERS:
            raise KnowledgeParseError("knowledge_extracted_text_too_large")
        units.append(ExtractionUnit(text=text, provenance={"page": page_index}))
    _validate_total_characters(units)
    return units


def _validate_xlsx_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_XLSX_ZIP_ENTRIES:
                raise KnowledgeParseError("knowledge_xlsx_too_many_entries")
            expanded_size = sum(info.file_size for info in infos)
            if expanded_size > MAX_XLSX_EXPANDED_BYTES:
                raise KnowledgeParseError("knowledge_xlsx_expanded_too_large")
    except zipfile.BadZipFile as exc:
        raise KnowledgeParseError("knowledge_xlsx_invalid") from exc


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float)):
        return str(value).strip()
    return str(value).strip()


def parse_xlsx(data: bytes) -> list[ExtractionUnit]:
    _validate_xlsx_archive(data)
    try:
        workbook = load_workbook(
            filename=BytesIO(data),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise KnowledgeParseError("knowledge_xlsx_invalid") from exc

    units: list[ExtractionUnit] = []
    extracted_total = 0
    try:
        for worksheet in workbook.worksheets:
            sheet_title = worksheet.title[:200]
            row_lines: list[str] = []
            row_start: int | None = None
            row_end: int | None = None
            current_chars = 0

            def flush(*, sheet_name: str = sheet_title) -> None:
                nonlocal row_lines, row_start, row_end, current_chars
                if not row_lines or row_start is None or row_end is None:
                    return
                text = normalize_extracted_text("\n".join(row_lines))
                if text:
                    units.append(
                        ExtractionUnit(
                            text=text,
                            provenance={
                                "sheet": sheet_name,
                                "row_start": row_start,
                                "row_end": row_end,
                            },
                        )
                    )
                row_lines = []
                row_start = None
                row_end = None
                current_chars = 0

            for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                values = [_cell_text(value) for value in cast(tuple[object, ...], row)]
                while values and not values[-1]:
                    values.pop()
                if not any(values):
                    continue
                line = f"row {row_index}: " + " | ".join(values)
                if row_lines and current_chars + len(line) + 1 > _XLSX_UNIT_TARGET_CHARACTERS:
                    flush()
                if row_start is None:
                    row_start = row_index
                row_end = row_index
                row_lines.append(line)
                current_chars += len(line) + 1
                extracted_total += len(line)
                if extracted_total > MAX_EXTRACTED_CHARACTERS:
                    raise KnowledgeParseError("knowledge_extracted_text_too_large")
            flush()
    finally:
        workbook.close()

    _validate_total_characters(units)
    return units


def parse_source(data: bytes, media_type: str) -> list[ExtractionUnit]:
    if media_type == PDF_MEDIA_TYPE:
        return parse_pdf(data)
    if media_type == XLSX_MEDIA_TYPE:
        return parse_xlsx(data)
    raise KnowledgeParseError("knowledge_media_type_unsupported")


def _find_chunk_boundary(text: str, start: int, hard_end: int) -> int:
    if hard_end >= len(text):
        return len(text)
    floor = start + CHUNK_MAX_CHARACTERS // 2
    for separator in ("\n\n", "\n", ". ", " "):
        boundary = text.rfind(separator, floor, hard_end)
        if boundary >= floor:
            return boundary + len(separator)
    return hard_end


def chunk_units(units: list[ExtractionUnit]) -> list[ChunkDraft]:
    chunks: list[ChunkDraft] = []
    for unit in units:
        text = normalize_extracted_text(unit.text)
        if not text:
            continue
        start = 0
        while start < len(text):
            hard_end = min(len(text), start + CHUNK_MAX_CHARACTERS)
            end = _find_chunk_boundary(text, start, hard_end)
            if end <= start:
                end = hard_end
            content = text[start:end].strip()
            if content:
                chunks.append(ChunkDraft(content=content, provenance=dict(unit.provenance)))
                if len(chunks) > MAX_SOURCE_CHUNKS:
                    raise KnowledgeParseError("knowledge_too_many_chunks")
            if end >= len(text):
                break
            next_start = max(start + 1, end - CHUNK_OVERLAP_CHARACTERS)
            start = next_start
    if not chunks:
        raise KnowledgeParseError("knowledge_no_extractable_text")
    return chunks
